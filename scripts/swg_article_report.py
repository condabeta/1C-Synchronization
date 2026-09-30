#!/usr/bin/env python
"""Turn the site's short SWG articles into SWG's own eleven-character codes.

    python scripts/swg_article_report.py

Анна asked on 30.09.2026 whether the articles of this brand could be made
eleven-digit automatically rather than by hand, the way the price list writes
them. They can, and no matching is needed: SWG writes the same number, prefixed
with "00-" and padded to eight digits. The site's 045488 is their 00-00045488.

The rule finds 1,145 of the 1,379 SWG products on the site. The remaining 234
are not in the September price list at all - discontinued, as Анна said of
002033 - and those are listed separately rather than converted, because a code
that is not in the price list is worse than the short number it replaced.

This also settles the articles this brand shares with Salux and Точка Зрения.
Once an SWG product carries 00-00002325, no other supplier's 002325 can be
confused with it.

Reads the site CSV exports rather than the site itself.
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from staging.db import db_session, fetch_all

SITE_DIR = Path(r"D:\projects\1C")
PRODUCTS_CSV = SITE_DIR / "oc_product.csv"
DESCRIPTIONS_CSV = SITE_DIR / "oc_product_description.csv"
MANUFACTURERS_CSV = SITE_DIR / "oc_manufacturer.csv"
OUTPUT = SITE_DIR / "SWG - замена артикулов на сайте.xlsx"

SITE_MANUFACTURER = "SWG"
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")


def swg_code(article: str) -> str | None:
    """The price list's spelling of a site article: 045488 -> 00-00045488."""
    return "00-" + article.zfill(8) if article.isdigit() else None


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Convert site SWG articles to the supplier's own")
    parser.add_argument("--out", type=Path, default=OUTPUT)
    return parser.parse_args()


def read_site() -> list[dict]:
    csv.field_size_limit(10_000_000)
    manufacturers, names = {}, {}
    with MANUFACTURERS_CSV.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            manufacturers[row["manufacturer_id"]] = row.get("name") or ""
    with DESCRIPTIONS_CSV.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            names.setdefault(row["product_id"], row.get("name") or "")

    products = []
    with PRODUCTS_CSV.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            if manufacturers.get(row.get("manufacturer_id") or "") != SITE_MANUFACTURER:
                continue
            products.append({
                "product_id": row.get("product_id"),
                "sku": (row.get("sku") or "").strip(),
                "name": names.get(row.get("product_id") or "", ""),
                "price": row.get("price"),
                "quantity": row.get("quantity"),
                "status": "в продаже" if (row.get("status") or "") == "1" else "выключен",
            })
    return products


def write_sheet(sheet, headers: list[str], widths: list[int], rows: list[list]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        sheet.append(row)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    sheet.freeze_panes = "A2"


def build(out_path: Path) -> int:
    with db_session() as conn:
        feed = {
            row["supplier_sku"]: row
            for row in fetch_all(
                conn,
                """
                SELECT sp.supplier_sku, sp.name, sp.price_retail, sp.is_available
                FROM supplier_products sp JOIN suppliers s ON s.id = sp.supplier_id
                WHERE s.code = 'swg'
                """,
            )
        }

    convert, missing = [], []
    for product in sorted(read_site(), key=lambda p: p["sku"]):
        code = swg_code(product["sku"])
        supplier = feed.get(code) if code else None
        if supplier:
            convert.append([
                product["sku"], code, product["name"], supplier["name"],
                product["price"], supplier["price_retail"], product["status"],
                product["product_id"],
            ])
        else:
            missing.append([
                product["sku"], code or "(артикул не числовой)", product["name"],
                product["price"], product["status"], product["product_id"],
            ])

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Замена артикулов"
    write_sheet(
        sheet,
        ["Артикул на сайте", "Артикул SWG (заменить на него)", "Наименование на сайте",
         "Наименование в прайсе SWG", "Цена на сайте, ₽", "Цена по прайсу, ₽", "Статус",
         "ID на сайте"],
        [16, 26, 50, 50, 15, 15, 11, 11],
        convert,
    )
    write_sheet(
        workbook.create_sheet("Нет в прайсе"),
        ["Артикул на сайте", "Ожидаемый артикул SWG", "Наименование на сайте",
         "Цена на сайте, ₽", "Статус", "ID на сайте"],
        [16, 24, 56, 15, 11, 11],
        missing,
    )
    workbook.save(out_path)

    total = len(convert) + len(missing)
    print(f"Товаров SWG на сайте: {total:,}")
    print(f"  артикул можно заменить автоматически: {len(convert):,} ({len(convert)/total*100:.1f}%)")
    print(f"  нет в сентябрьском прайсе SWG:        {len(missing):,}")
    print(f"\nПравило: «00-» + артикул, дополненный нулями до восьми знаков.")
    print(f"Файл: {out_path}")
    return 0


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)
        sys.stderr.reconfigure(encoding="utf-8")
    args = parse_args()
    for path in (PRODUCTS_CSV, DESCRIPTIONS_CSV, MANUFACTURERS_CSV):
        if not path.exists():
            print(f"Выгрузка сайта не найдена: {path}", file=sys.stderr)
            return 1
    return build(args.out)


if __name__ == "__main__":
    raise SystemExit(main())
