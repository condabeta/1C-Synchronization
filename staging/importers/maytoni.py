"""Maytoni's own price list - seven files, seven brands.

Until now Maytoni goods came only through Dekomo. Анна asked for them to be
separated, because the client buys Maytoni direct and the direct price differs,
and pointed at https://shared.maytoni.ru/files/PRICE_RRC/ - a folder holding a
file per brand: Maytoni, Technical, Outdoor, Ledstrip, Lighting control, Freya
and Voltega.

The files are worth more than their prices. Each row carries a **barcode**,
which only three of our suppliers give us at all, and the workbook has a
**photograph embedded against every product** - 7,735 of them - which matters
because Dekomo's image links are dead and Arlight's are blocked.

The seven files do not share a layout. Some name the brand in a column and some
only in the file, the price is «РРЦ, руб.» or «РРЦ» or «РРЦ, шт», and only some
carry «Полное наименование». So every sheet is read through its own header row
rather than by fixed positions, and the missing pieces are filled from what is
there: a name built from the type, series and colour when none is given, a
brand taken from the file name.

Section titles sit between the products, filling only the first column. They
are skipped by requiring a price.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable, Iterator

import openpyxl

from staging.importers.common import clean, content_hash, parse_decimal

SUPPLIER_CODE = "maytoni"
SOURCE_CODE = "price_xlsx"
PRICE_DIR = Path(r"D:\projects\1C\Майтони")
PRICE_URL = "https://shared.maytoni.ru/files/PRICE_RRC/"

ARTICLE_HEADER = "Артикул"
# In order of preference: a per-piece price beats a per-metre one, and the tape
# file carries both.
PRICE_HEADERS = ("РРЦ, руб.", "РРЦ, шт", "РРЦ", "РРЦ, м")
NAME_HEADERS = ("Полное наименование", "Наименование")
BRAND_HEADER = "Бренд"
BARCODE_HEADERS = ("Штрихкод", "Штрих-код")
PHOTO_HEADER = "Фото"

# Columns that describe the product rather than identify it. Kept as attributes
# so a product page can show them.
ATTR_SKIP = {ARTICLE_HEADER, BRAND_HEADER, PHOTO_HEADER, *PRICE_HEADERS, *NAME_HEADERS, *BARCODE_HEADERS}

# "РРЦ_Technical_2026 NEW.xlsx" -> "Technical"
FILE_BRAND_RE = re.compile(r"РРЦ[_\s]+(.+?)[_\s]+\d{4}", re.I)


@dataclass
class SheetImage:
    row: int
    data: bytes
    extension: str


def brand_from_filename(path: Path) -> str | None:
    match = FILE_BRAND_RE.search(path.stem)
    return match.group(1).replace("_", " ").strip() if match else None


def _header_row(sheet) -> tuple[int, dict[str, int]] | None:
    """The row that names the columns, and where each one sits."""
    for index, row in enumerate(sheet.iter_rows(min_row=1, max_row=12, values_only=True), start=1):
        values = [clean(cell) for cell in row]
        if ARTICLE_HEADER in values:
            headers = {value: position for position, value in enumerate(values) if value}
            return index, headers
    return None


def read_images(path: Path) -> dict[str, dict[int, SheetImage]]:
    """{sheet: {row: image}} - one photograph per product row."""
    workbook = openpyxl.load_workbook(path)
    images: dict[str, dict[int, SheetImage]] = {}
    for name in workbook.sheetnames:
        sheet = workbook[name]
        found: dict[int, SheetImage] = {}
        for image in getattr(sheet, "_images", []):
            anchor = getattr(image.anchor, "_from", None)
            if anchor is None:
                continue
            found[anchor.row + 1] = SheetImage(
                row=anchor.row + 1,
                data=image._data(),
                extension=(image.format or "png").lower(),
            )
        images[name] = found
    workbook.close()
    return images


def _name_for(values: dict[str, Any], headers: dict[str, int], article: str) -> str:
    for header in NAME_HEADERS:
        text = clean(values.get(header))
        if text:
            return text
    # No name column: build one the way a person would read the row.
    parts = [clean(values.get(header)) for header in
             ("Тип светильника", "Тип продукта", "Серия", "Коллекция", "Цвет")]
    built = " ".join(part for part in parts if part)
    return built or article


def iter_maytoni_rows(
    directory: Path = PRICE_DIR,
    *,
    progress: Callable[[str], None] | None = None,
) -> Iterator[dict[str, Any]]:
    """Every priced article across the seven files, images included."""
    for path in sorted(directory.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        file_brand = brand_from_filename(path)
        images = read_images(path)
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        count = 0

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            header = _header_row(sheet)
            if not header:
                continue
            header_index, headers = header
            sheet_images = images.get(sheet_name, {})

            for row_number, row in enumerate(
                sheet.iter_rows(min_row=header_index + 1, values_only=True), start=header_index + 1
            ):
                values = {name: (row[position] if position < len(row) else None)
                          for name, position in headers.items()}
                article = clean(values.get(ARTICLE_HEADER))
                if not article or article == "<->":
                    continue
                price = next(
                    (parse_decimal(values.get(header)) for header in PRICE_HEADERS
                     if parse_decimal(values.get(header)) is not None),
                    None,
                )
                if price is None or price <= 0:
                    continue  # a section title, or an article sold on request

                brand = clean(values.get(BRAND_HEADER)) or file_brand
                barcode = next((clean(values.get(h)) for h in BARCODE_HEADERS if clean(values.get(h))), None)
                attrs = {
                    name: clean(values.get(name))
                    for name in headers
                    if name not in ATTR_SKIP and clean(values.get(name)) not in (None, "-")
                }
                attrs["file"] = path.name
                image = sheet_images.get(row_number)

                item = {
                    "supplier_sku": article[:128],
                    "supplier_sku_raw": article,
                    "name": _name_for(values, headers, article),
                    "brand": brand,
                    "manufacturer_code": article[:255],
                    "description": None,
                    "supplier_category": clean(values.get("Тип светильника")) or clean(values.get("Тип продукта")),
                    "supplier_category_path": " / ".join(
                        part for part in (brand, clean(values.get("Коллекция")), clean(values.get("Серия"))) if part
                    ),
                    "price": price,
                    "price_retail": price,  # РРЦ: the price Maytoni sets for the shelf
                    "price_old": None,
                    "stock_qty": None,  # the price list carries no stock
                    "is_available": 1,
                    "product_url": None,
                    "barcode": barcode[:64] if barcode else None,
                    "images_json": [],  # filled once the photo is written out
                    "attributes_json": attrs,
                    "raw_data_json": None,
                    "_image": image,
                    "_source_file": path.name,
                }
                item["content_hash"] = content_hash(
                    {
                        "supplier_sku": item["supplier_sku"],
                        "name": item["name"],
                        "brand": brand,
                        "price": str(price),
                        "barcode": barcode,
                        "attributes_json": attrs,
                    }
                )
                count += 1
                yield item
        workbook.close()
        if progress:
            with_photo = sum(len(v) for v in images.values())
            progress(f"  {path.name}: {count:,} товаров, {with_photo:,} фотографий")
