from __future__ import annotations

import colorsys
import hashlib
import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable, Iterator

import openpyxl
import pandas as pd
from pymysql.connections import Connection

from staging.importers.common import (
    ImportStats,
    clean,
    close_stale_runs,
    content_hash,
    count_supplier_products,
    file_sha256,
    finish_import_run,
    get_supplier_and_source,
    load_existing_hashes,
    log_import_error,
    parse_decimal,
    start_import_run,
    upsert_product_batch,
)

SUPPLIER_CODE = "viasvet"
SOURCE_CODE = "price_xlsx"
DEFAULT_XLSX_PATH = (
    r"D:\projects\1C\Виа Свет\ViaSvet_led_профиль_блоки_питания_лента_ПОСТУПЛЕНИЕ5 (3).xlsx"
)
DEFAULT_PHOTOS_DIR = r"D:\projects\1C\Виа Свет\на сайт"
BATCH_SIZE = 200

PROFILE_SHEETS = (
    "Профили для LED ленты",
    "Профили для светильников",
)
ACCESSORY_SHEET = "КОМПЛЕКТУЮЩИЕ"
TAPE_SHEET = "ЛЕНТА!!!"
PS_SHEETS = (
    "Блоки сетка 12V",
    "Блоки сетка 24V",
    "Ультратонкие 12V",
    "Блоки IP67 12v и 24v",
)
UNAVAILABLE_SKUS = {"SP264", "SP268", "SP284"}
DELETED_SKUS = {"SP284"}
HIDDEN_SKUS = {"SP268"}

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
SUFFIX_TOKENS = ("B2", "RGD", "GD", "CP", "GR", "W", "T", "B")


def _report(progress: Callable[[str], None] | None, message: str) -> None:
    if progress:
        progress(message)
    else:
        print(message, flush=True)


def load_sku_aliases(conn: Connection, supplier_id: int) -> dict[str, str]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT alias_sku, canonical_sku
            FROM sku_aliases
            WHERE supplier_id = %s OR supplier_id IS NULL
            """,
            (supplier_id,),
        )
        return {row["alias_sku"]: row["canonical_sku"] for row in cur.fetchall()}


def scan_photo_folders(photos_dir: Path) -> dict[str, list[str]]:
    index: dict[str, list[str]] = {}
    if not photos_dir.is_dir():
        return index

    skip_dirs = {"ВЫВЕЛИ или НЕТ в наличии"}
    for folder in sorted(photos_dir.iterdir()):
        if not folder.is_dir() or folder.name in skip_dirs:
            continue
        images = sorted(
            str(p.resolve())
            for p in folder.iterdir()
            if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS
        )
        if images:
            index[folder.name.upper()] = images
    return index


def _photo_lookup_keys(sku: str, aliases: dict[str, str]) -> list[str]:
    keys = [sku.upper()]
    canonical = aliases.get(sku) or aliases.get(sku.upper())
    if canonical:
        keys.append(canonical.upper())
    if "В" in sku:
        keys.append(sku.upper().replace("В", "B"))

    base = re.sub(r"-(?:NOCLIP|CLIPS|RRC\d+)$", "", sku.upper())
    if base != sku.upper():
        keys.append(base)
        if "В" in base:
            keys.append(base.replace("В", "B"))

    length_variant = re.match(r"^(SP\d+)\.\d+$", base)
    if length_variant:
        keys.append(length_variant.group(1))

    return list(dict.fromkeys(keys))


def resolve_images(
    sku: str,
    photo_index: dict[str, list[str]],
    aliases: dict[str, str],
) -> list[str]:
    for key in _photo_lookup_keys(sku, aliases):
        images = photo_index.get(key)
        if images:
            return images
    return []


def _extract_sp_suffix(compact_tail: str) -> str:
    tail = compact_tail.upper().replace("В", "B")
    for token in SUFFIX_TOKENS:
        if tail.startswith(token):
            return token
    return ""


def extract_profile_sku(name: str, kit: str | None, rrc: Any) -> tuple[str, str] | None:
    text = clean(name)
    if not text:
        return None

    first_token = re.sub(r"\s+", "", text.split()[0])
    match = re.match(r"^(SP\d+(?:\.\d+)?)([A-ZА-Я]{0,4})?$", first_token, re.I)
    if not match:
        return None

    base = match.group(1).upper()
    suffix = _extract_sp_suffix(match.group(2) or "")
    if not suffix and "Черный" in text and "черный экран" not in text.lower():
        suffix = "B"

    sku = base + suffix
    raw = text.split()[0]

    # Check if SKU should be skipped (unavailable/deleted/hidden)
    if any(unavail in sku for unavail in UNAVAILABLE_SKUS):
        return None

    variant_tags: list[str] = []
    name_lower = text.lower()
    kit_lower = (kit or "").lower()
    if "без крепеж" in name_lower:
        variant_tags.append("NOCLIP")
    elif kit and "3 крепеж" in kit_lower and "2 метра" not in kit_lower:
        variant_tags.append("CLIPS")

    if variant_tags:
        sku = f"{sku}-{'-'.join(variant_tags)}"
    elif rrc is not None:
        rrc_text = clean(rrc)
        if rrc_text and sku in {"SP259W", "SP280"}:
            sku = f"{sku}-RRC{rrc_text.replace('.', '')}"

    return sku, raw


def extract_accessory_sku(name: str) -> tuple[str, str] | None:
    text = clean(name)
    if not text:
        return None
    compact = re.sub(r"\s+", "", text)
    match = re.match(r"^((?:PS|DO|EN|CL)SP[A-Z0-9]+)", compact, re.I)
    if not match:
        return None
    sku = match.group(1).upper()
    return sku, sku


def extract_tape_sku(row: pd.Series) -> tuple[str, str] | None:
    sku = clean(row.iloc[1] if len(row) > 1 else None)
    if not sku or not sku.upper().startswith("VST-"):
        return None
    return sku.upper(), sku


def _tape_length_variant(description: str | None) -> str:
    if not description:
        return ""
    match = re.search(r"Длина:\s*(\d+)\s*м", description, re.I)
    if match:
        return f"-{match.group(1)}M"
    return ""


def extract_ps_sku(row: pd.Series, sheet: str) -> tuple[str, str, str] | None:
    if sheet == "Блоки IP67 12v и 24v":
        sku = clean(row.iloc[1] if len(row) > 1 else None)
        name = clean(row.iloc[0] if len(row) else None)
    else:
        sku = clean(row.iloc[2] if len(row) > 2 else None)
        name = clean(row.iloc[1] if len(row) > 1 else None)

    if not sku or sku.lower() == "артикул":
        return None
    if not re.match(r"^[A-Z0-9]", sku, re.I):
        return None
    sku = re.sub(r"\s+", "", sku).upper()
    if not name:
        name = sku
    return sku, sku, name


def _cell(row: pd.Series, index: int) -> Any:
    return row.iloc[index] if len(row) > index else None


def _has_price(value: Any) -> bool:
    if value is None:
        return False
    text = str(value).strip().lower()
    return text not in ("", "nan", "none") and "наличии" not in text


@dataclass(frozen=True)
class StockBand:
    """A stock range read from a cell colour. ViaSvet gives ranges, not counts."""

    label: str
    # The lowest quantity the band promises, stored as stock_qty so "in stock"
    # checks work. It is a floor, never a count - the label is the truth.
    floor: int
    available: bool


# ViaSvet encodes stock as the background colour of one column, with a legend at
# the top of the sheet: green > 1000, yellow 500-1000, red 100-500.
BAND_GREEN = StockBand(">1000", 1000, True)
BAND_YELLOW = StockBand("500-1000", 500, True)
BAND_RED = StockBand("100-500", 100, True)
BAND_OUT = StockBand("нет в наличии", 0, False)


def _band_for_fill(cell: Any) -> StockBand | None:
    """Classify a cell fill by hue rather than by exact colour.

    The cells were painted by hand and do not match their own legend: the legend
    green is 00B050 while the painted green is 1F7012, and yellow appears as both
    F0EA00 and FFFF00. Matching exact RGB would miss 44 of 47 rows on the LED
    profile sheet, so the hue decides.
    """
    fill = cell.fill
    if fill is None or fill.fill_type in (None, "none"):
        return None
    colour = fill.start_color
    if colour is None or colour.type != "rgb":
        return None  # theme colours here are the white page background
    rgb = str(colour.rgb or "")
    if len(rgb) != 8:
        return None
    red, green, blue = (int(rgb[i : i + 2], 16) / 255 for i in (2, 4, 6))
    hue, saturation, value = colorsys.rgb_to_hsv(red, green, blue)
    if saturation < 0.35 or value < 0.2:
        return None  # white, grey or near-black - not a stock marker
    degrees = hue * 360
    if degrees < 20 or degrees >= 340:
        return BAND_RED
    if degrees < 70:
        return BAND_YELLOW
    if degrees < 170:
        return BAND_GREEN
    return None


# A stock column is one where at least this many data cells carry a band colour.
STOCK_COLUMN_MIN_CELLS = 3
STOCK_SCAN_COLUMNS = 20
FIRST_DATA_ROW = 5


def _find_stock_column(sheet: Any) -> int | None:
    """The column holding the stock colours, found rather than hardcoded.

    It moves between sheets - column H on the profile sheets, M on the power
    supply sheets - and the "Наличие->" header does not sit above it, so the
    column with the most band-coloured cells is taken.
    """
    best_column, best_count = None, 0
    for column in range(1, min(sheet.max_column, STOCK_SCAN_COLUMNS) + 1):
        count = sum(
            1
            for row in range(FIRST_DATA_ROW, sheet.max_row + 1)
            if _band_for_fill(sheet.cell(row, column))
        )
        if count > best_count:
            best_column, best_count = column, count
    return best_column if best_count >= STOCK_COLUMN_MIN_CELLS else None


def read_stock_bands(xlsx_path: Path) -> dict[tuple[str, int], StockBand]:
    """{(sheet, 1-based Excel row): band} for every row with a stock marker."""
    book = openpyxl.load_workbook(xlsx_path)
    bands: dict[tuple[str, int], StockBand] = {}
    for sheet in book.worksheets:
        column = _find_stock_column(sheet)
        if column is None:
            continue
        for row in range(FIRST_DATA_ROW, sheet.max_row + 1):
            cell = sheet.cell(row, column)
            if "нет в наличии" in str(cell.value or "").lower():
                bands[(sheet.title, row)] = BAND_OUT
                continue
            band = _band_for_fill(cell)
            if band:
                bands[(sheet.title, row)] = band
    return bands


def normalize_product(
    *,
    sheet: str,
    sku: str,
    sku_raw: str,
    name: str,
    price: Any,
    price_retail: Any,
    kit: str | None = None,
    packaging: str | None = None,
    extra_attrs: dict[str, Any] | None = None,
    photo_index: dict[str, list[str]],
    aliases: dict[str, str],
    stock_band: StockBand | None = None,
    store_raw: bool = False,
    raw_row: dict[str, Any] | None = None,
) -> dict[str, Any] | None:
    parsed_price = parse_decimal(price)
    if parsed_price is None:
        return None

    parsed_rrc = parse_decimal(price_retail)
    images = resolve_images(sku, photo_index, aliases)

    attrs: dict[str, Any] = {"sheet": sheet}
    if kit:
        attrs["kit_description"] = kit
    if packaging:
        attrs["packaging"] = packaging
    if extra_attrs:
        attrs.update(extra_attrs)
    if stock_band:
        attrs["stock_band"] = stock_band.label

    normalized = {
        "supplier_sku": sku,
        "supplier_sku_raw": sku_raw,
        "name": name,
        "brand": "ViaSvet",
        "manufacturer_code": sku_raw,
        "supplier_category": sheet,
        "supplier_category_path": sheet,
        "price": parsed_price,
        "price_retail": parsed_rrc or parsed_price,
        "price_old": parsed_rrc,
        # No colour means the sheet says nothing about this row, so stock stays
        # unknown rather than being read as zero.
        "stock_qty": Decimal(stock_band.floor) if stock_band else None,
        "is_available": (1 if stock_band.available else 0) if stock_band else 1,
        "product_url": "https://www.viasvet.ru",
        "barcode": None,
        "images_json": images,
        "attributes_json": attrs,
        "raw_data_json": raw_row if store_raw else None,
    }
    normalized["content_hash"] = content_hash(
        {
            "supplier_sku": sku,
            "name": name,
            "price": str(parsed_price),
            "price_retail": str(parsed_rrc) if parsed_rrc is not None else None,
            "stock_band": stock_band.label if stock_band else None,
            "images_json": images,
            "attributes_json": attrs,
        }
    )
    return normalized


def parse_profile_rows(
    xlsx_path: Path,
    photo_index: dict[str, list[str]],
    aliases: dict[str, str],
    *,
    bands: dict[tuple[str, int], StockBand] | None = None,
    store_raw: bool = False,
) -> Iterator[dict[str, Any]]:
    for sheet in PROFILE_SHEETS:
        df = pd.read_excel(xlsx_path, sheet_name=sheet, header=None)
        for index, row in df.iterrows():
            name = clean(_cell(row, 0))
            if not name or "SP" not in name.upper():
                continue
            sku_pair = extract_profile_sku(name, clean(_cell(row, 3)), _cell(row, 5))
            if not sku_pair:
                continue
            sku, sku_raw = sku_pair
            item = normalize_product(
                sheet=sheet,
                sku=sku,
                sku_raw=sku_raw,
                name=name,
                price=_cell(row, 6),
                price_retail=_cell(row, 5),
                kit=clean(_cell(row, 3)),
                packaging=clean(_cell(row, 4)),
                photo_index=photo_index,
                aliases=aliases,
                stock_band=(bands or {}).get((sheet, index + 1)),
                store_raw=store_raw,
                raw_row=row.to_dict(),
            )
            if item:
                yield item


def parse_accessory_rows(
    xlsx_path: Path,
    photo_index: dict[str, list[str]],
    aliases: dict[str, str],
    *,
    bands: dict[tuple[str, int], StockBand] | None = None,
    store_raw: bool = False,
) -> Iterator[dict[str, Any]]:
    df = pd.read_excel(xlsx_path, sheet_name=ACCESSORY_SHEET, header=None)
    for index, row in df.iterrows():
        name = clean(_cell(row, 0))
        sku_pair = extract_accessory_sku(name or "")
        if not sku_pair:
            continue
        sku, sku_raw = sku_pair
        item = normalize_product(
            sheet=ACCESSORY_SHEET,
            sku=sku,
            sku_raw=sku_raw,
            name=name or sku,
            price=_cell(row, 6),
            price_retail=_cell(row, 5),
            kit=clean(_cell(row, 3)),
            photo_index=photo_index,
            aliases=aliases,
            stock_band=(bands or {}).get((ACCESSORY_SHEET, index + 1)),
            store_raw=store_raw,
            raw_row=row.to_dict(),
        )
        if item:
            yield item


def parse_tape_rows(
    xlsx_path: Path,
    photo_index: dict[str, list[str]],
    aliases: dict[str, str],
    *,
    bands: dict[tuple[str, int], StockBand] | None = None,
    store_raw: bool = False,
) -> Iterator[dict[str, Any]]:
    df = pd.read_excel(xlsx_path, sheet_name=TAPE_SHEET, header=None)
    for index, row in df.iterrows():
        sku_pair = extract_tape_sku(row)
        if not sku_pair:
            continue
        sku, sku_raw = sku_pair
        description = clean(_cell(row, 3))
        length_variant = _tape_length_variant(description)
        sku = f"{sku}{length_variant}"
        name = clean(_cell(row, 2)) or sku_raw
        price = _cell(row, 6) if _has_price(_cell(row, 6)) else _cell(row, 5)
        item = normalize_product(
            sheet=TAPE_SHEET,
            sku=sku,
            sku_raw=sku_raw,
            name=name,
            price=price,
            price_retail=_cell(row, 7),
            kit=clean(_cell(row, 3)),
            extra_attrs={"price_per_meter": clean(_cell(row, 5))},
            photo_index=photo_index,
            aliases=aliases,
            stock_band=(bands or {}).get((TAPE_SHEET, index + 1)),
            store_raw=store_raw,
            raw_row=row.to_dict(),
        )
        if item:
            yield item


def parse_ps_rows(
    xlsx_path: Path,
    photo_index: dict[str, list[str]],
    aliases: dict[str, str],
    *,
    bands: dict[tuple[str, int], StockBand] | None = None,
    store_raw: bool = False,
) -> Iterator[dict[str, Any]]:
    for sheet in PS_SHEETS:
        df = pd.read_excel(xlsx_path, sheet_name=sheet, header=None)
        for index, row in df.iterrows():
            parsed = extract_ps_sku(row, sheet)
            if not parsed:
                continue
            sku, sku_raw, name = parsed
            price = _cell(row, 11) if _has_price(_cell(row, 11)) else _cell(row, 10)
            if not _has_price(price):
                continue
            attrs = {
                "current_a": clean(_cell(row, 4 if sheet == "Блоки IP67 12v и 24v" else 3)),
                "voltage": clean(_cell(row, 5 if sheet == "Блоки IP67 12v и 24v" else 4)),
                "power_w": clean(_cell(row, 6 if sheet == "Блоки IP67 12v и 24v" else 5)),
                "dimensions": clean(_cell(row, 7 if sheet == "Блоки IP67 12v и 24v" else 6)),
                "ip_rating": clean(_cell(row, 9 if sheet == "Блоки IP67 12v и 24v" else 8)),
            }
            item = normalize_product(
                sheet=sheet,
                sku=sku,
                sku_raw=sku_raw,
                name=name,
                price=price,
                price_retail=_cell(row, 10),
                extra_attrs={k: v for k, v in attrs.items() if v},
                photo_index=photo_index,
                aliases=aliases,
                stock_band=(bands or {}).get((sheet, index + 1)),
                store_raw=store_raw,
                raw_row=row.to_dict(),
            )
            if item:
                yield item


def iter_viasvet_products(
    xlsx_path: Path,
    photos_dir: Path,
    aliases: dict[str, str] | None = None,
    *,
    profiles_only: bool = False,
    store_raw: bool = False,
) -> Iterator[dict[str, Any]]:
    photo_index = scan_photo_folders(photos_dir)
    alias_map = aliases or {}
    # Stock lives in cell colours, which pandas cannot see, so it is read once
    # with openpyxl and looked up by (sheet, row) as each parser walks its rows.
    bands = read_stock_bands(xlsx_path)
    common = {"bands": bands, "store_raw": store_raw}

    yield from parse_profile_rows(xlsx_path, photo_index, alias_map, **common)
    if profiles_only:
        return
    yield from parse_accessory_rows(xlsx_path, photo_index, alias_map, **common)
    yield from parse_tape_rows(xlsx_path, photo_index, alias_map, **common)
    yield from parse_ps_rows(xlsx_path, photo_index, alias_map, **common)


def count_viasvet_products(
    xlsx_path: Path,
    photos_dir: Path,
    aliases: dict[str, str] | None = None,
    *,
    profiles_only: bool = False,
) -> int:
    return sum(
        1
        for _ in iter_viasvet_products(
            xlsx_path,
            photos_dir,
            aliases,
            profiles_only=profiles_only,
        )
    )


def combined_source_hash(xlsx_path: Path, photos_dir: Path) -> str:
    digest = hashlib.sha256()
    digest.update(file_sha256(xlsx_path).encode("utf-8"))
    if photos_dir.is_dir():
        for folder in sorted(photos_dir.iterdir()):
            if folder.is_dir():
                digest.update(folder.name.encode("utf-8"))
                for image in sorted(folder.iterdir()):
                    if image.is_file() and image.suffix.lower() in IMAGE_EXTENSIONS:
                        digest.update(str(image.stat().st_size).encode("utf-8"))
                        digest.update(str(int(image.stat().st_mtime)).encode("utf-8"))
    return digest.hexdigest()


def import_viasvet_xlsx(
    conn: Connection,
    xlsx_path: Path,
    photos_dir: Path,
    limit: int | None = None,
    skip_rows: int = 0,
    profiles_only: bool = False,
    progress: Callable[[str], None] | None = None,
    store_raw: bool = False,
) -> ImportStats:
    if not xlsx_path.exists():
        raise FileNotFoundError(f"ViaSvet XLSX not found: {xlsx_path}")

    supplier_id, source_id = get_supplier_and_source(conn, SUPPLIER_CODE, SOURCE_CODE)
    aliases = load_sku_aliases(conn, supplier_id)

    _report(progress, f"Reading ViaSvet Excel ({xlsx_path.name})...")
    total_expected = count_viasvet_products(
        xlsx_path,
        photos_dir,
        aliases,
        profiles_only=profiles_only,
    )
    scope = "profiles only" if profiles_only else "all sheets"
    _report(progress, f"Found ~{total_expected:,} products ({scope}).")

    close_stale_runs(conn, supplier_id)
    source_hash = combined_source_hash(xlsx_path, photos_dir)
    source_label = f"{xlsx_path}|{photos_dir}"
    run_id = start_import_run(conn, supplier_id, source_id, source_label, source_hash)
    conn.commit()

    stats = ImportStats()
    batch: list[dict[str, Any]] = []
    existing_hashes: dict[str, str] = {}
    batch_number = 0
    product_index = 0

    try:
        for item in iter_viasvet_products(
            xlsx_path,
            photos_dir,
            aliases,
            profiles_only=profiles_only,
            store_raw=store_raw,
        ):
            stats.rows_total += 1
            try:
                product_index += 1
                if product_index <= skip_rows:
                    continue
                if limit is not None and product_index - skip_rows > limit:
                    break

                batch.append(item)
                if len(batch) >= BATCH_SIZE:
                    batch_number += 1
                    skus = [row["supplier_sku"] for row in batch]
                    existing_hashes.update(load_existing_hashes(conn, supplier_id, skus))
                    upsert_product_batch(conn, supplier_id, run_id, batch, existing_hashes, stats)
                    conn.commit()
                    batch.clear()
                    _report(
                        progress,
                        (
                            f"  batch {batch_number}: products {product_index:,} | "
                            f"db {count_supplier_products(conn, supplier_id):,} | "
                            f"new {stats.rows_imported:,} | updated {stats.rows_updated:,} | "
                            f"unchanged {stats.rows_skipped:,} | errors {stats.rows_errors:,}"
                        ),
                    )
            except Exception as exc:
                stats.rows_errors += 1
                stats.errors.append({"row": product_index, "error": str(exc), "sku": item.get("supplier_sku")})
                log_import_error(
                    conn,
                    run_id,
                    product_index,
                    item.get("supplier_sku"),
                    str(exc),
                    item,
                )
                conn.commit()

        if batch:
            batch_number += 1
            skus = [row["supplier_sku"] for row in batch]
            existing_hashes.update(load_existing_hashes(conn, supplier_id, skus))
            upsert_product_batch(conn, supplier_id, run_id, batch, existing_hashes, stats)
            conn.commit()
            _report(
                progress,
                (
                    f"  batch {batch_number}: products {product_index:,} | "
                    f"db {count_supplier_products(conn, supplier_id):,} | "
                    f"new {stats.rows_imported:,} | updated {stats.rows_updated:,} | "
                    f"unchanged {stats.rows_skipped:,} | errors {stats.rows_errors:,}"
                ),
            )

        status = "success"
        if stats.rows_errors and (stats.rows_imported or stats.rows_updated):
            status = "partial"
        elif stats.rows_errors and not (stats.rows_imported or stats.rows_updated):
            status = "failed"

        error_summary = None
        if stats.errors:
            error_summary = "; ".join(
                f"{item.get('sku', '?')}: {item['error']}" for item in stats.errors[:5]
            )

        finish_import_run(conn, run_id, stats, status, error_summary)
        conn.commit()
        _report(progress, "Import completed.")
        return stats
    except Exception:
        conn.rollback()
        finish_import_run(conn, run_id, stats, "failed", "Import aborted due to unexpected error")
        conn.commit()
        raise
